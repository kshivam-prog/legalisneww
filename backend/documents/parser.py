"""
Document parsing module for PDF, DOCX, and text files
"""

import os
from typing import Tuple, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parse different document formats to text"""
    
    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """
        Parse PDF file and extract text
        
        Args:
            file_path: Path to PDF file
            
        Returns:
            Extracted text content
            
        Raises:
            ImportError: If PyPDF2 is not installed
            Exception: If PDF cannot be parsed
        """
        try:
            import PyPDF2
        except ImportError:
            raise ImportError("PyPDF2 is required for PDF parsing. Install with: pip install PyPDF2")
        
        text_content = []
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)
                
                logger.info(f"Parsing PDF with {num_pages} pages: {file_path}")
                
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    text = page.extract_text()
                    if text:
                        text_content.append(f"--- Page {page_num} ---\n{text}")
                    
                    # Log progress for large files
                    if page_num % 10 == 0:
                        logger.info(f"Parsed {page_num}/{num_pages} pages")
            
            result = "\n\n".join(text_content)
            logger.info(f"Successfully parsed PDF: {len(result)} characters extracted")
            return result
            
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {str(e)}")
            raise
    
    @staticmethod
    def parse_docx(file_path: str) -> str:
        """
        Parse DOCX file and extract text
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Extracted text content
            
        Raises:
            ImportError: If python-docx is not installed
            Exception: If DOCX cannot be parsed
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")
        
        text_content = []
        
        try:
            doc = Document(file_path)
            logger.info(f"Parsing DOCX with {len(doc.paragraphs)} paragraphs: {file_path}")
            
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)
            
            # Extract text from tables
            for table_idx, table in enumerate(doc.tables, 1):
                text_content.append(f"\n--- Table {table_idx} ---\n")
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    text_content.append(row_text)
            
            result = "\n".join(text_content)
            logger.info(f"Successfully parsed DOCX: {len(result)} characters extracted")
            return result
            
        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {str(e)}")
            raise
    
    @staticmethod
    def parse_text(content: str) -> str:
        """
        Validate and clean plain text content
        
        Args:
            content: Text content
            
        Returns:
            Cleaned text content
        """
        if not content or not isinstance(content, str):
            raise ValueError("Content must be non-empty string")
        
        # Remove excessive whitespace
        text = "\n".join(line.strip() for line in content.split("\n") if line.strip())
        
        logger.info(f"Parsed plain text: {len(text)} characters")
        return text
    
    @staticmethod
    def parse_file(file_path: str) -> Tuple[str, str]:
        """
        Auto-detect and parse file based on extension
        
        Args:
            file_path: Path to file
            
        Returns:
            Tuple of (text_content, file_type)
            
        Raises:
            ValueError: If file format is not supported
            FileNotFoundError: If file doesn't exist
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = Path(file_path).suffix.lower()
        
        logger.info(f"Parsing file with extension {file_ext}: {file_path}")
        
        if file_ext == '.pdf':
            content = DocumentParser.parse_pdf(file_path)
            return content, 'pdf'
        elif file_ext == '.docx':
            content = DocumentParser.parse_docx(file_path)
            return content, 'docx'
        elif file_ext in ['.txt', '.text']:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            return DocumentParser.parse_text(content), 'text'
        else:
            raise ValueError(f"Unsupported file format: {file_ext}. Supported: .pdf, .docx, .txt")


def parse_document(file_path: str) -> Tuple[str, str]:
    """
    Convenience function to parse a document
    
    Args:
        file_path: Path to document file
        
    Returns:
        Tuple of (text_content, file_type)
    """
    return DocumentParser.parse_file(file_path)
